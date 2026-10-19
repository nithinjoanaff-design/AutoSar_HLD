import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header(doc, title_text, subtitle_text):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run(title_text)
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run(subtitle_text)
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(100, 110, 120)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(25, 55, 109)
    return p

# -------------------------------------------------------------
# 1. GENERATE ACADEMIC INTEGRITY DECLARATION (.docx)
# -------------------------------------------------------------
doc1 = docx.Document()
for s in doc1.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.85)
    s.right_margin = Inches(0.85)

add_header(doc1, "Academic Integrity Declaration", 
           "AUTOSAR HLD Intelligence & Traceability Assistant  |  Student declaration | Editable signature form")

add_section_heading(doc1, "Student and Project Details")

fields1 = [
    ("University", "SASTRA Deemed University (confirm if required)"),
    ("Student name", "____________________________________________"),
    ("Register number", "____________________________________________"),
    ("Programme / Department", "____________________________________________"),
    ("Project title", "AUTOSAR HLD Intelligence & Traceability Assistant"),
    ("Project guide", "____________________________________________"),
    ("Date", "____________________________________________"),
]

t1 = doc1.add_table(rows=len(fields1) + 1, cols=2)
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t1, color="B0C4DE")

# Header row
t1.rows[0].cells[0].text = "Field"
t1.rows[0].cells[1].text = "Student to Complete"
for i in range(2):
    c = t1.rows[0].cells[i]
    set_cell_shading(c, "102C57")
    set_cell_margins(c, top=100, bottom=100)
    p = c.paragraphs[0]
    p.runs[0].font.name = 'Calibri'
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

for r_idx, (fld, val) in enumerate(fields1, start=1):
    c0, c1 = t1.rows[r_idx].cells[0], t1.rows[r_idx].cells[1]
    c0.text = fld
    c1.text = val
    c0.width = Inches(2.3)
    c1.width = Inches(4.5)
    set_cell_margins(c0, top=70, bottom=70)
    set_cell_margins(c1, top=70, bottom=70)
    if r_idx % 2 == 1:
        set_cell_shading(c0, "F5F7FA")
        set_cell_shading(c1, "F5F7FA")
    for cell in (c0, c1):
        cell.paragraphs[0].runs[0].font.name = 'Calibri'
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
    c0.paragraphs[0].runs[0].font.bold = True

add_section_heading(doc1, "Declaration")

decl_paras = [
    "I declare that the capstone submission titled “AUTOSAR HLD Intelligence & Traceability Assistant” is an honest account of the work completed by me and my team, as applicable. I have accurately represented the work performed, individual contributions, implementation status, evaluation results, and known limitations. I have not knowingly fabricated experimental results or presented generated examples as real-world engineering evidence.",
    "I will cite external code, datasets, models, documents, and other sources used in the submission. I will disclose AI tools and pre-trained models used, including tools used for coding assistance, in the separate AI-Tool Usage Declaration. I understand that AI-assisted output must be reviewed and verified by me, and that I remain responsible for the correctness and explainability of the submitted work.",
    "I understand that the evaluator may ask me to explain the system design, implementation, prompts, configuration, tests, results, and any AI-generated or externally sourced artifacts included in the submission."
]

for dp in decl_paras:
    p = doc1.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(dp)
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(40, 40, 40)

# Callout box
t_box = doc1.add_table(rows=1, cols=1)
t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
c_box = t_box.rows[0].cells[0]
c_box.width = Inches(6.8)
set_cell_shading(c_box, "EBF3FA")
set_cell_margins(c_box, top=100, bottom=100, left=150, right=150)
set_table_borders(t_box, color="4682B4", sz="8")
p_box = c_box.paragraphs[0]
p_box.paragraph_format.space_before = Pt(0)
p_box.paragraph_format.space_after = Pt(0)
r_bold = p_box.add_run("Before signing: ")
r_bold.font.bold = True
r_bold.font.color.rgb = RGBColor(16, 44, 87)
r_text = p_box.add_run("Read the declaration carefully. Keep only statements that are true for your work. Complete all personal fields, sign and date the form, and follow any official institution-provided declaration wording if it differs.")
for r in (r_bold, r_text):
    r.font.name = 'Calibri'
    r.font.size = Pt(9.5)

add_section_heading(doc1, "Student Confirmation")

p_sign = doc1.add_paragraph()
p_sign.paragraph_format.space_before = Pt(10)
p_sign.paragraph_format.space_after = Pt(6)
p_sign.paragraph_format.line_spacing = 1.3
r_sign = p_sign.add_run("Student name (print): ____________________________________    Signature: __________________________    Date: ______________")
r_sign.font.name = 'Calibri'
r_sign.font.size = Pt(10)

p_note = doc1.add_paragraph()
p_note.paragraph_format.space_before = Pt(4)
p_note.paragraph_format.space_after = Pt(0)
r_note = p_note.add_run("For team submissions, every student must complete and sign a declaration, as required by the instructor.")
r_note.font.name = 'Calibri'
r_note.font.size = Pt(9)
r_note.font.italic = True
r_note.font.color.rgb = RGBColor(100, 100, 100)

doc1.save("d:/autosar-hld-assistant/Academic_Integrity_Declaration.docx")
print("Saved Academic_Integrity_Declaration.docx successfully.")


# -------------------------------------------------------------
# 2. GENERATE AI-TOOL USAGE DECLARATION (.docx)
# -------------------------------------------------------------
doc2 = docx.Document()
for s in doc2.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.85)
    s.right_margin = Inches(0.85)

add_header(doc2, "AI-Tool Usage Declaration", 
           "AUTOSAR HLD Intelligence & Traceability Assistant  |  Tool disclosure table | Editable student verification form")

add_section_heading(doc2, "Student and Project Details")

fields2 = [
    ("University", "SASTRA Deemed University (confirm if required)"),
    ("Student name", "____________________________________________"),
    ("Register number", "____________________________________________"),
    ("Project title", "AUTOSAR HLD Intelligence & Traceability Assistant"),
    ("Date", "____________________________________________"),
]

t2 = doc2.add_table(rows=len(fields2) + 1, cols=2)
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t2, color="B0C4DE")

t2.rows[0].cells[0].text = "Field"
t2.rows[0].cells[1].text = "Student to Complete"
for i in range(2):
    c = t2.rows[0].cells[i]
    set_cell_shading(c, "102C57")
    set_cell_margins(c, top=90, bottom=90)
    p = c.paragraphs[0]
    p.runs[0].font.name = 'Calibri'
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

for r_idx, (fld, val) in enumerate(fields2, start=1):
    c0, c1 = t2.rows[r_idx].cells[0], t2.rows[r_idx].cells[1]
    c0.text = fld
    c1.text = val
    c0.width = Inches(2.2)
    c1.width = Inches(4.6)
    set_cell_margins(c0, top=60, bottom=60)
    set_cell_margins(c1, top=60, bottom=60)
    if r_idx % 2 == 1:
        set_cell_shading(c0, "F5F7FA")
        set_cell_shading(c1, "F5F7FA")
    for cell in (c0, c1):
        cell.paragraphs[0].runs[0].font.name = 'Calibri'
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
    c0.paragraphs[0].runs[0].font.bold = True

add_section_heading(doc2, "Tool-Use Disclosure")

p_tool_desc = doc2.add_paragraph()
p_tool_desc.paragraph_format.space_before = Pt(2)
p_tool_desc.paragraph_format.space_after = Pt(6)
r_td = p_tool_desc.add_run(
    "List every AI tool and pre-trained model used in the project, including tools used for coding assistance. "
    "The rows below are populated based on the exact implementation of the AUTOSAR HLD Intelligence & Traceability Assistant. "
    "Verify each row, update any details specific to your environment, and record your verification notes before signing."
)
r_td.font.name = 'Calibri'
r_td.font.size = Pt(9.5)
r_td.font.color.rgb = RGBColor(50, 50, 50)

# Exact table requested with student's models (Ollama deepseek-r1:1.5b / qwen2.5, all-MiniLM-L6-v2, Antigravity, etc.)
tools_data = [
    (
        "Ollama + deepseek-r1:1.5b\n(or Qwen2.5:3B)",
        "Local language-model inference for RAG responses; generating Executive Summary, Architectural Breakdown, and verified citations.",
        "Generated answers in the AUTOSAR HLD Assistant Q&A interface.",
        "Verified answers against uploaded HLD PDFs; tested reasoning stripping (<think> blocks) and ground-truth page numbers. Verified with pytest tests/test_rag_pipeline.py."
    ),
    (
        "Sentence Transformers\n(all-MiniLM-L6-v2) &\nCross-Encoder\n(ms-marco-MiniLM-L-6-v2)",
        "384-dimensional dense vector embeddings for semantic retrieval and Cross-Encoder deep query-passage re-ranking.",
        "FAISS vector store, similarity search, candidate passage re-ranking.",
        "Verified vector dimensions (384), cosine similarity scoring, and reranking precision across 10 evaluation test questions."
    ),
    (
        "ChatGPT / OpenAI API\n(where applicable)",
        "Architecture conceptualization, comparative testing, prompt refinement, and evaluation dataset drafting.",
        "Prompt templates in src/llm_provider.py and evaluation questions in evaluate.py.",
        "Manually reviewed technical automotive accuracy (AUTOSAR SWC, RTE, BSW terminology) and executed pytest evaluation."
    ),
    (
        "Antigravity Coding Assistant",
        "Assisted in code structuring, test generation, cross-encoder integration, Streamlit UI components, and packaging automation.",
        "Source modules (src/), test suite (tests/), start_app.bat, stop_app.bat, app.py.",
        "Reviewed all assisted source code; executed full test suite achieving 78/78 passed tests; verified manual end-to-end execution."
    ),
    (
        "Other AI Tools / Pretrained Models",
        "Specify any additional models (e.g. OpenRouter free tier models meta-llama/llama-3.3-70b-instruct).",
        "Artifact affected: ________________________",
        "Student verification notes: ________________________________________________"
    )
]

t_tools = doc2.add_table(rows=len(tools_data) + 1, cols=4)
t_tools.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_tools, color="B0C4DE")

headers = ["AI Model / Tool", "Purpose", "Artifact Affected", "Verified by Student"]
widths = [Inches(1.5), Inches(1.8), Inches(1.6), Inches(1.9)]

for col_idx, text in enumerate(headers):
    c = t_tools.rows[0].cells[col_idx]
    c.text = text
    c.width = widths[col_idx]
    set_cell_shading(c, "102C57")
    set_cell_margins(c, top=100, bottom=100, left=100, right=100)
    p = c.paragraphs[0]
    p.runs[0].font.name = 'Calibri'
    p.runs[0].font.size = Pt(9.5)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

for row_idx, row_data in enumerate(tools_data, start=1):
    for col_idx, text in enumerate(row_data):
        c = t_tools.rows[row_idx].cells[col_idx]
        c.text = text
        c.width = widths[col_idx]
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        if row_idx % 2 == 1:
            set_cell_shading(c, "F5F7FA")
        p = c.paragraphs[0]
        p.runs[0].font.name = 'Calibri'
        p.runs[0].font.size = Pt(8.5)
        p.paragraph_format.line_spacing = 1.15
        if col_idx == 0:
            p.runs[0].font.bold = True

add_section_heading(doc2, "Student Attestation")

p_att = doc2.add_paragraph()
p_att.paragraph_format.space_before = Pt(3)
p_att.paragraph_format.space_after = Pt(4)
p_att.paragraph_format.line_spacing = 1.15
r_att = p_att.add_run(
    "I confirm that the disclosure above is complete and accurate to the best of my knowledge. "
    "I have reviewed the AI-assisted or model-generated outputs used in the submission and remain responsible for their correctness. "
    "I have not represented unverified AI output as confirmed fact and have not fabricated results. "
    "I can explain the submitted implementation and the role played by each declared tool/model."
)
r_att.font.name = 'Calibri'
r_att.font.size = Pt(9.5)
r_att.font.color.rgb = RGBColor(40, 40, 40)

# Student Action Warning Box
t_box2 = doc2.add_table(rows=1, cols=1)
t_box2.alignment = WD_TABLE_ALIGNMENT.CENTER
c_box2 = t_box2.rows[0].cells[0]
c_box2.width = Inches(6.8)
set_cell_shading(c_box2, "EBF3FA")
set_cell_margins(c_box2, top=80, bottom=80, left=120, right=120)
set_table_borders(t_box2, color="4682B4", sz="8")
p_box2 = c_box2.paragraphs[0]
p_box2.paragraph_format.space_before = Pt(0)
p_box2.paragraph_format.space_after = Pt(0)
r_b2_bold = p_box2.add_run("Required student action: ")
r_b2_bold.font.bold = True
r_b2_bold.font.color.rgb = RGBColor(16, 44, 87)
r_b2_text = p_box2.add_run("Edit the draft tool rows to match your exact tools and environment. Verify how each tool was tested, sign the form, and include it in your Declarations/ submission folder.")
for r in (r_b2_bold, r_b2_text):
    r.font.name = 'Calibri'
    r.font.size = Pt(9)

add_section_heading(doc2, "Signature")

t_sig = doc2.add_table(rows=2, cols=3)
t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_sig, color="B0C4DE")

sig_headers = ["Student Name (print)", "Signature", "Date"]
for idx, text in enumerate(sig_headers):
    c = t_sig.rows[0].cells[idx]
    c.text = text
    set_cell_shading(c, "102C57")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    p = c.paragraphs[0]
    p.runs[0].font.name = 'Calibri'
    p.runs[0].font.size = Pt(9.5)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

sig_vals = ["__________________________", "__________________________", "__________________"]
for idx, text in enumerate(sig_vals):
    c = t_sig.rows[1].cells[idx]
    c.text = text
    set_cell_margins(c, top=140, bottom=140, left=100, right=100)
    p = c.paragraphs[0]
    p.runs[0].font.name = 'Calibri'
    p.runs[0].font.size = Pt(9.5)

p_inst_note = doc2.add_paragraph()
p_inst_note.paragraph_format.space_before = Pt(6)
p_inst_note.paragraph_format.space_after = Pt(0)
r_in = p_inst_note.add_run(
    "The instructor’s instruction PDF requires every student to sign the AI-tool usage declaration and shows these columns: "
    "AI Model/Tool, Purpose, Artifact affected, Verified by student. It does not provide an official declaration template or exact wording."
)
r_in.font.name = 'Calibri'
r_in.font.size = Pt(8.5)
r_in.font.italic = True
r_in.font.color.rgb = RGBColor(110, 110, 110)

doc2.save("d:/autosar-hld-assistant/AI_Tool_Usage_Declaration.docx")
print("Saved AI_Tool_Usage_Declaration.docx successfully.")
